from __future__ import annotations

import os
import smtplib
import ssl
import threading
import time
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path
from string import Template
from typing import List, Union

import docx
from openpyxl import load_workbook


class EmailSender:
    @staticmethod
    def _sender_email_target(from_email, to_email, subject, text, server):
        message = MIMEMultipart("alternative")
        message["Subject"] = subject
        message["From"] = from_email
        message["To"] = to_email

        # Turn these into plain/html MIMEText objects
        part1 = MIMEText(text, "plain")
        # part2 = MIMEText(html, "html")

        # Add HTML/plain-text parts to MIMEMultipart message
        # The email client will try to render the last part first
        message.attach(part1)
        # message.attach(part2)
        tries = 0
        while tries < 3:
            try:
                server.sendmail(from_email, to_email, message.as_string())
                print("Email sended", flush=True)
                break
            except Exception:
                # TODO: handle the specific exception
                print("Email send failed, retrying in 3 seconds.", flush=True)
                tries += 1
            time.sleep(3)

    @staticmethod
    def send(from_email, to_email, subject, text, server):
        thr = threading.Thread(
            target=EmailSender._sender_email_target,
            args=(from_email, to_email, subject, text, server),
        )
        thr.run()

    @staticmethod
    def _sender_from_template_target(template, to_email, **data):
        sended = False
        tries = 0
        while not sended and tries < 3:
            try:
                context = ssl.create_default_context()

                url = "smtp.gmail.com"
                port = 465
                username = os.getenv("EMAIL_USERNAME")
                password = os.getenv("EMAIL_PASSWORD")

                with open(f"/src/templates/{template}.txt", "r") as f:
                    text_template = Template(f.read())

                with smtplib.SMTP_SSL(url, port, context=context) as server:
                    server.login(username, password)
                    text = text_template.substitute(**data)
                    sender_email = username
                    subject = text.partition("\n")[0]
                    message = MIMEMultipart("alternative")
                    message["Subject"] = subject
                    message["From"] = sender_email
                    message["To"] = to_email

                    # Turn these into plain/html MIMEText objects
                    part1 = MIMEText(text, "plain")
                    # part2 = MIMEText(html, "html")

                    # Add HTML/plain-text parts to MIMEMultipart message
                    # The email client will try to render the last part first
                    message.attach(part1)
                    # message.attach(part2)
                    server.sendmail(sender_email, to_email, message.as_string())
                    print("Email sended", flush=True)
                    sended = True
            except Exception:
                # TODO: handle the specific exception
                print("Email send failed, retrying in 3 seconds.", flush=True)
                tries += 1

            if not sended:
                time.sleep(3)

    @staticmethod
    def send_from_template(template, to_email, **data):
        thr = threading.Thread(
            target=EmailSender._sender_from_template_target,
            args=(template, to_email),
            kwargs=data,
        )
        thr.run()


def send_email(from_email, to_email, subject, text, server):
    EmailSender.send(from_email, to_email, subject, text, server)


def send_from_template(template, to_email, **data):
    EmailSender.send_from_template(template, to_email, **data)


def clean_temp_after_run(func):
    def wrapper(*args, **kwargs):
        val = func(*args, **kwargs)
        for temp_file in Path().glob("_temp.*"):
            os.remove(temp_file)
        return val

    return wrapper


def create_temp(fd, ext):
    temp = open("_tmp." + ext, "wb")
    fd = fd.read()
    temp.write(fd)
    return temp


@clean_temp_after_run
def check_file(file, document: str) -> bool:
    if not file:
        return False
    d = "".join(document.split("-"))
    secure_key = d[: len(d) - 5].lower()
    ext = file.name[-4:]
    if ext == "docx":
        doc = docx.Document(file)
        if doc.core_properties.keywords == secure_key:
            return True
    elif ext == "xlsx":
        wb = load_workbook(file)
        if wb.properties.keywords == secure_key:
            return True
    return False
